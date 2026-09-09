# Copyright (c) 2026, abdulloh and contributors
# For license information, please see license.txt

import frappe
from frappe import _
from frappe.utils import flt


CATEGORY_MAP = {
    "Покупатели": "customer",
    "Поставщики": "supplier",
    "Расходы": "expense",
    "Дивиденды": "dividend",
    "Сотрудники": "employee",
    "Перемещения": "transfer",
}

CATEGORY_LABELS = {
    "customer": "Покупатели",
    "supplier": "Поставщики",
    "expense": "Расходы",
    "dividend": "Дивиденды",
    "employee": "Сотрудники",
    "transfer": "Перемещения",
    "other": "Прочие",
}


def execute(filters=None):
    columns = get_columns()
    data, expense_summaries, employee_group_summaries = get_data(filters)
    summary_html = get_summary_html(data, expense_summaries, employee_group_summaries)
    return columns, data, summary_html


def get_columns():
    return [
        {"fieldname": "posting_date", "label": _("Сана"), "fieldtype": "Date", "width": 100},
        {"fieldname": "account", "label": _("Касса счёт"), "fieldtype": "Link", "options": "Account", "width": 180},
        {"fieldname": "direction", "label": _("Кирим/Чиқим"), "fieldtype": "Data", "width": 100},
        {"fieldname": "description", "label": _("Категория"), "fieldtype": "Data", "width": 250},
        {"fieldname": "schet", "label": _("Счёт"), "fieldtype": "Data", "width": 140},
        {"fieldname": "tip1", "label": _("Тип 1"), "fieldtype": "Data", "width": 130},
        {"fieldname": "summa", "label": _("Сумма"), "fieldtype": "Currency", "options": "currency", "width": 130},
        {"fieldname": "currency", "label": _("Валюта"), "fieldtype": "Link", "options": "Currency", "hidden": 1},
        {"fieldname": "remarks", "label": _("Изоҳ"), "fieldtype": "Data", "width": 200},
        {"fieldname": "voucher_type", "label": _("Тип"), "fieldtype": "Data", "width": 0, "hidden": 1},
        {"fieldname": "voucher_no", "label": _("Документ"), "fieldtype": "Dynamic Link", "options": "voucher_type", "width": 160},
    ]


def get_employee_group_map():
    """Return {employee_id: group_name}. Har employee ko'pi bilan 1 guruhda."""
    rows = frappe.get_all(
        "Employee Group Table",
        fields=["employee", "parent"],
    )
    return {r.employee: r.parent for r in rows if r.employee}


def get_data(filters):
    cash_accounts = get_cash_accounts(filters)
    if not cash_accounts:
        return [], {}, {}

    opening_balance = get_opening_balance(cash_accounts, filters)
    transactions = get_transactions(cash_accounts, filters)

    pe_vouchers = [r.voucher_no for r in transactions if r.voucher_type == "Payment Entry"]
    je_vouchers = [r.voucher_no for r in transactions if r.voucher_type == "Journal Entry"]

    pe_info = get_payment_entry_info_batch(pe_vouchers)
    je_info = get_journal_entry_info_batch(je_vouchers)
    je_remarks = get_journal_entry_remarks_batch(je_vouchers)

    data = []

    # Filterlar
    filter_party_type = filters.get("party_type")
    filter_party = filters.get("party")
    category_filter_val = filters.get("category")
    filter_category = CATEGORY_MAP.get(category_filter_val)

    expense_summaries = {}
    employee_group_summaries = {}
    employee_group_map = get_employee_group_map()
    balance = opening_balance
    total_kirim = 0
    total_chiqim = 0

    for row in transactions:
        kirim = flt(row.debit_in_account_currency)
        chiqim = flt(row.credit_in_account_currency)

        info = resolve_transaction_info(row, pe_info, je_info, cash_accounts)

        # Начальный остаток (Temporary Opening) — Прочие emas, opening balansga qo'sh
        if info["category"] == "opening":
            opening_balance += kirim - chiqim
            balance += kirim - chiqim
            continue

        # Category filter
        if filter_category and info["category"] != filter_category:
            balance += kirim - chiqim
            continue

        # Party filter
        if filter_party_type and info.get("party_type") != filter_party_type:
            balance += kirim - chiqim
            continue
        if filter_party and info.get("party") != filter_party:
            balance += kirim - chiqim
            continue

        balance += kirim - chiqim
        total_kirim += kirim
        total_chiqim += chiqim

        # Xarajatlarni guruhlash
        if info["category"] == "expense":
            desc = info["description"]
            if desc not in expense_summaries:
                expense_summaries[desc] = {"kirim": 0, "chiqim": 0}
            expense_summaries[desc]["kirim"] += kirim
            expense_summaries[desc]["chiqim"] += chiqim

        # Xodimlarni employee group bo'yicha guruhlash
        emp_group = ""
        if info["category"] == "employee":
            emp_group = employee_group_map.get(info.get("party"), "Без группы")
            if emp_group not in employee_group_summaries:
                employee_group_summaries[emp_group] = {"kirim": 0, "chiqim": 0}
            employee_group_summaries[emp_group]["kirim"] += kirim
            employee_group_summaries[emp_group]["chiqim"] += chiqim

        data.append({
            "posting_date": row.posting_date,
            "account": row.account,
            "direction": "Кирим" if kirim else "Чиқим",
            "employee_group": emp_group,
            "description": strip_category_prefix(info["description"]),
            "category": info["category"],
            "schet": info.get("schet", ""),
            "tip1": info.get("tip1", ""),
            "summa": kirim if kirim else chiqim,
            "remarks": get_remarks(row, pe_info, je_remarks),
            "voucher_type": row.voucher_type,
            "voucher_no": row.voucher_no,
            "currency": row.account_currency,
            "kirim": kirim,
            "chiqim": chiqim,
        })

    final_data = list(data)

    # opening_balance va closing balance ni summary HTML uchun saqlash
    if final_data:
        final_data[0]["_opening_balance"] = opening_balance
        final_data[-1]["_closing_balance"] = balance

    return final_data, expense_summaries, employee_group_summaries


def get_cash_accounts(filters):
    conditions = {}
    if filters.get("mode_of_payment"):
        conditions["parent"] = filters["mode_of_payment"]

    accounts = frappe.get_all(
        "Mode of Payment Account",
        filters=conditions,
        fields=["default_account"],
        pluck="default_account"
    )

    return list(set(a for a in accounts if a))


def get_opening_balance(cash_accounts, filters):
    placeholders = ", ".join(["%s"] * len(cash_accounts))

    result = frappe.db.sql("""
        SELECT IFNULL(SUM(debit_in_account_currency) - SUM(credit_in_account_currency), 0)
        FROM `tabGL Entry`
        WHERE account IN ({placeholders})
          AND posting_date < %s
          AND is_cancelled = 0
    """.format(placeholders=placeholders),
        tuple(cash_accounts) + (filters["from_date"],)
    )

    return flt(result[0][0]) if result else 0


def get_transactions(cash_accounts, filters):
    placeholders = ", ".join(["%s"] * len(cash_accounts))

    return frappe.db.sql("""
        SELECT
            posting_date, voucher_type, voucher_no,
            party_type, party, against,
            debit_in_account_currency, credit_in_account_currency,
            account, account_currency
        FROM `tabGL Entry`
        WHERE account IN ({placeholders})
          AND posting_date BETWEEN %s AND %s
          AND is_cancelled = 0
        ORDER BY posting_date, creation
    """.format(placeholders=placeholders),
        tuple(cash_accounts) + (filters["from_date"], filters["to_date"]),
        as_dict=True
    )


def get_payment_entry_info_batch(voucher_nos):
    if not voucher_nos:
        return {}

    entries = frappe.db.sql("""
        SELECT name, party_type, party, payment_type, remarks
        FROM `tabPayment Entry`
        WHERE name IN %s
    """, (voucher_nos,), as_dict=True)

    return {e.name: e for e in entries}


def get_journal_entry_info_batch(voucher_nos):
    if not voucher_nos:
        return {}

    entries = frappe.db.sql("""
        SELECT jea.parent, jea.account, jea.party_type, jea.party,
               acc.root_type, acc.account_type, acc.account_name,
               pacc.account_name AS parent_account_name
        FROM `tabJournal Entry Account` jea
        LEFT JOIN `tabAccount` acc ON acc.name = jea.account
        LEFT JOIN `tabAccount` pacc ON pacc.name = acc.parent_account
        WHERE jea.parent IN %s
    """, (voucher_nos,), as_dict=True)

    result = {}
    for e in entries:
        result.setdefault(e.parent, []).append(e)
    return result


def get_journal_entry_remarks_batch(voucher_nos):
    if not voucher_nos:
        return {}

    entries = frappe.db.sql("""
        SELECT name, user_remark
        FROM `tabJournal Entry`
        WHERE name IN %s
    """, (voucher_nos,), as_dict=True)

    return {e.name: (e.user_remark or "") for e in entries}


def get_remarks(row, pe_info, je_remarks):
    if row.voucher_type == "Payment Entry" and row.voucher_no in pe_info:
        return pe_info[row.voucher_no].get("remarks") or ""
    if row.voucher_type == "Journal Entry" and row.voucher_no in je_remarks:
        return je_remarks[row.voucher_no] or ""
    return ""


def strip_category_prefix(desc):
    for prefix in ("Расходы: ", "Дивиденды: "):
        if desc.startswith(prefix):
            return desc[len(prefix):]
    return desc


def resolve_transaction_info(row, pe_info, je_info, cash_accounts):
    # 1. GL Entry'da party bor
    if row.party_type and row.party:
        party_name = get_party_name(row.party_type, row.party)
        display_name = party_name or row.party
        suffix = "Приход" if flt(row.debit_in_account_currency) > 0 else "Расход"
        return {
            "description": f"{display_name} ({suffix})",
            "category": get_category_from_party_type(row.party_type),
            "party_type": row.party_type,
            "party": row.party,
        }

    # 2. Payment Entry
    if row.voucher_type == "Payment Entry" and row.voucher_no in pe_info:
        pe = pe_info[row.voucher_no]
        if pe.payment_type == "Internal Transfer":
            return {"description": "Перемещение", "category": "transfer", "party_type": None, "party": None}
        if pe.party_type and pe.party:
            party_name = get_party_name(pe.party_type, pe.party)
            display_name = party_name or pe.party
            suffix = "Приход" if pe.payment_type == "Receive" else "Расход"
            return {
                "description": f"{display_name} ({suffix})",
                "category": get_category_from_party_type(pe.party_type),
                "party_type": pe.party_type,
                "party": pe.party,
            }

    # 3. Journal Entry
    if row.voucher_type == "Journal Entry" and row.voucher_no in je_info:
        for acc in je_info[row.voucher_no]:
            if acc.account in cash_accounts:
                continue
            if acc.account_type == "Temporary":
                return {"description": "Начальный остаток", "category": "opening", "party_type": None, "party": None}
            if acc.party_type and acc.party:
                party_name = get_party_name(acc.party_type, acc.party)
                return {
                    "description": party_name or acc.party,
                    "category": get_category_from_party_type(acc.party_type),
                    "party_type": acc.party_type,
                    "party": acc.party,
                }
            if acc.root_type == "Expense":
                # Xarajat schoti Kassa Rasxod qatoridagi "Тип 1", uning ota schoti esa "Счёт".
                # Ikkalasi ham hujjatga alohida yozilmaydi — schotlar daraxtidan tiklanadi.
                return {
                    "description": f"Расходы: {acc.account_name}",
                    "category": "expense",
                    "party_type": None,
                    "party": None,
                    "schet": acc.parent_account_name or "",
                    "tip1": acc.account_name or "",
                }
            if acc.root_type == "Equity":
                return {"description": f"Дивиденды: {acc.account_name}", "category": "dividend", "party_type": None, "party": None}

    # 4. Against field (fallback)
    if row.against:
        against_account = row.against.split(",")[0].strip() if "," in row.against else row.against

        is_cash = frappe.db.get_value("Mode of Payment Account", {"default_account": against_account}, "parent")
        if is_cash:
            direction = "из" if flt(row.debit_in_account_currency) > 0 else "в"
            return {"description": f"Перемещение {direction} {is_cash}", "category": "transfer", "party_type": None, "party": None}

        acc_info = frappe.db.get_value(
            "Account", against_account,
            ["account_name", "root_type", "account_type", "parent_account"], as_dict=True
        )
        if acc_info:
            if acc_info.account_type == "Temporary":
                return {"description": "Начальный остаток", "category": "opening", "party_type": None, "party": None}
            if acc_info.root_type == "Expense":
                parent_name = (
                    frappe.get_cached_value("Account", acc_info.parent_account, "account_name")
                    if acc_info.parent_account else ""
                ) or ""
                return {
                    "description": f"Расходы: {acc_info.account_name}",
                    "category": "expense",
                    "party_type": None,
                    "party": None,
                    "schet": parent_name,
                    "tip1": acc_info.account_name or "",
                }
            if acc_info.root_type == "Equity":
                return {"description": f"Дивиденды: {acc_info.account_name}", "category": "dividend", "party_type": None, "party": None}
            if acc_info.account_type == "Receivable":
                return {"description": acc_info.account_name, "category": "customer", "party_type": "Customer", "party": None}
            if acc_info.account_type == "Payable":
                return {"description": acc_info.account_name, "category": "supplier", "party_type": "Supplier", "party": None}
            return {"description": acc_info.account_name, "category": "other", "party_type": None, "party": None}

    return {"description": row.voucher_no or "", "category": "other", "party_type": None, "party": None}


def get_category_from_party_type(party_type):
    return {"Customer": "customer", "Supplier": "supplier", "Employee": "employee"}.get(party_type, "other")


def get_party_name(party_type, party):
    field = {"Customer": "customer_name", "Supplier": "supplier_name", "Employee": "employee_name"}.get(party_type)
    if field:
        return frappe.db.get_value(party_type, party, field)
    return party


def compute_summary(data, expense_summaries=None, employee_group_summaries=None):
    """Category bo'yicha kirim/chiqim yig'indilari + opening/closing. HTML va Excel export uchun umumiy."""
    cats = {c: {"kirim": 0, "chiqim": 0} for c in
            ("customer", "supplier", "expense", "dividend", "transfer", "employee", "other")}

    opening = 0
    closing_balance = 0
    for row in data:
        if "_opening_balance" in row:
            opening = flt(row["_opening_balance"])
        if "_closing_balance" in row:
            closing_balance = flt(row["_closing_balance"])

    for row in data:
        if row.get("is_total"):
            continue
        category = row.get("category") or "other"
        if category not in cats:
            category = "other"
        cats[category]["kirim"] += flt(row.get("kirim"))
        cats[category]["chiqim"] += flt(row.get("chiqim"))

    total_kirim = sum(c["kirim"] for c in cats.values())
    total_chiqim = sum(c["chiqim"] for c in cats.values())
    closing = opening + total_kirim - total_chiqim
    if closing_balance:
        closing = closing_balance

    return {
        "opening": opening,
        "closing": closing,
        "cats": cats,
        "expense_sub": expense_summaries or {},
        "employee_sub": employee_group_summaries or {},
    }


def get_summary_html(data, expense_summaries=None, employee_group_summaries=None):
    if not data:
        return ""

    s = compute_summary(data, expense_summaries, employee_group_summaries)
    opening = s["opening"]
    closing = s["closing"]
    customer_kirim, customer_chiqim = s["cats"]["customer"]["kirim"], s["cats"]["customer"]["chiqim"]
    supplier_kirim, supplier_chiqim = s["cats"]["supplier"]["kirim"], s["cats"]["supplier"]["chiqim"]
    expense_kirim, expense_chiqim = s["cats"]["expense"]["kirim"], s["cats"]["expense"]["chiqim"]
    dividend_kirim, dividend_chiqim = s["cats"]["dividend"]["kirim"], s["cats"]["dividend"]["chiqim"]
    transfer_kirim, transfer_chiqim = s["cats"]["transfer"]["kirim"], s["cats"]["transfer"]["chiqim"]
    employee_kirim, employee_chiqim = s["cats"]["employee"]["kirim"], s["cats"]["employee"]["chiqim"]
    other_kirim, other_chiqim = s["cats"]["other"]["kirim"], s["cats"]["other"]["chiqim"]

    def fmt(val):
        return f"{flt(val):,.2f}"

    def dash_or_val(val):
        return "—" if flt(val) == 0 else f"<span style='color: inherit;'>{fmt(val)}</span>"

    # Расходы subcategory qatorlarini tayyorlash
    expense_sub_rows = ""
    if expense_summaries:
        for desc, totals in expense_summaries.items():
            display_name = desc.replace("Расходы: ", "") if desc.startswith("Расходы: ") else desc
            sub_kirim = fmt(totals["kirim"]) if totals["kirim"] else "—"
            sub_chiqim = fmt(totals["chiqim"]) if totals["chiqim"] else "—"
            expense_sub_rows += f"""
                <tr class="dds-expense-sub" style="display: none; background-color: #fff8e1;">
                    <td style="padding: 8px 10px 8px 30px; border: 1px solid #ddd; font-style: italic;">{display_name}</td>
                    <td style="padding: 8px 10px; border: 1px solid #ddd; text-align: right; color: #388e3c;">{sub_kirim}</td>
                    <td style="padding: 8px 10px; border: 1px solid #ddd; text-align: right; color: #d32f2f;">{sub_chiqim}</td>
                </tr>"""

    # Сотрудники ostidagi employee group qatorlari (bosilganda ochiladi)
    employee_group_sub_rows = ""
    if employee_group_summaries:
        for group in sorted(employee_group_summaries, key=lambda g: (g == "Без группы", g)):
            totals = employee_group_summaries[group]
            grp_kirim = fmt(totals["kirim"]) if totals["kirim"] else "—"
            grp_chiqim = fmt(totals["chiqim"]) if totals["chiqim"] else "—"
            employee_group_sub_rows += f"""
                <tr class="dds-employee-sub" style="display: none; background-color: #f5f5f5;">
                    <td style="padding: 8px 10px 8px 30px; border: 1px solid #ddd; font-style: italic;">{group}</td>
                    <td style="padding: 8px 10px; border: 1px solid #ddd; text-align: right; color: #388e3c;">{grp_kirim}</td>
                    <td style="padding: 8px 10px; border: 1px solid #ddd; text-align: right; color: #d32f2f;">{grp_chiqim}</td>
                </tr>"""

    expense_arrow = '<span id="dds-expense-arrow" style="margin-right: 5px; font-size: 10px;">&#9654;</span>' if expense_summaries else ""
    expense_cursor = "cursor: pointer;" if expense_summaries else ""
    expense_onclick = """onclick="(function(){
        var rows = document.querySelectorAll('.dds-expense-sub');
        var arrow = document.getElementById('dds-expense-arrow');
        if (!rows.length) return;
        var visible = rows[0].style.display !== 'none';
        for (var i = 0; i < rows.length; i++) { rows[i].style.display = visible ? 'none' : 'table-row'; }
        arrow.innerHTML = visible ? '&#9654;' : '&#9660;';
    })()" """ if expense_summaries else ""

    employee_arrow = '<span id="dds-employee-arrow" style="margin-right: 5px; font-size: 10px;">&#9654;</span>' if employee_group_summaries else ""
    employee_cursor = "cursor: pointer;" if employee_group_summaries else ""
    employee_onclick = """onclick="(function(){
        var rows = document.querySelectorAll('.dds-employee-sub');
        var arrow = document.getElementById('dds-employee-arrow');
        if (!rows.length) return;
        var visible = rows[0].style.display !== 'none';
        for (var i = 0; i < rows.length; i++) { rows[i].style.display = visible ? 'none' : 'table-row'; }
        arrow.innerHTML = visible ? '&#9654;' : '&#9660;';
    })()" """ if employee_group_summaries else ""

    html = f"""
    <div style="margin-top: 20px; padding: 15px; background-color: #f9f9f9; border-radius: 5px;">
        <table style="width: 100%; border-collapse: collapse; background: white;">
            <thead>
                <tr style="background-color: #f0f0f0;">
                    <th style="padding: 10px; text-align: left; border: 1px solid #ddd; width: 40%;"></th>
                    <th style="padding: 10px; text-align: right; border: 1px solid #ddd; width: 30%; color: #388e3c; font-weight: bold;">Кирим</th>
                    <th style="padding: 10px; text-align: right; border: 1px solid #ddd; width: 30%; color: #d32f2f; font-weight: bold;">Чиқим</th>
                </tr>
            </thead>
            <tbody>
                <tr style="background-color: #e3f2fd;">
                    <td style="padding: 10px; border: 1px solid #ddd; font-weight: bold;">Начальный остаток</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; font-weight: bold;" colspan="2">{fmt(opening)}</td>
                </tr>
                <tr>
                    <td style="padding: 10px; border: 1px solid #ddd;">Покупатели</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #388e3c;">{fmt(customer_kirim) if customer_kirim else '—'}</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #d32f2f;">{fmt(customer_chiqim) if customer_chiqim else '—'}</td>
                </tr>
                <tr style="background-color: #fafafa;">
                    <td style="padding: 10px; border: 1px solid #ddd;">Поставщики</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #388e3c;">{fmt(supplier_kirim) if supplier_kirim else '—'}</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #d32f2f;">{fmt(supplier_chiqim) if supplier_chiqim else '—'}</td>
                </tr>
                <tr>
                    <td style="padding: 10px; border: 1px solid #ddd;">Дивиденды</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #388e3c;">{fmt(dividend_kirim) if dividend_kirim else '—'}</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #d32f2f;">{fmt(dividend_chiqim) if dividend_chiqim else '—'}</td>
                </tr>
                <tr style="background-color: #fafafa; {employee_cursor}" {employee_onclick}>
                    <td style="padding: 10px; border: 1px solid #ddd;">{employee_arrow}Сотрудники</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #388e3c;">{fmt(employee_kirim) if employee_kirim else '—'}</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #d32f2f;">{fmt(employee_chiqim) if employee_chiqim else '—'}</td>
                </tr>
                {employee_group_sub_rows}
                <tr>
                    <td style="padding: 10px; border: 1px solid #ddd;">Перемещения</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #388e3c;">{fmt(transfer_kirim) if transfer_kirim else '—'}</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #d32f2f;">{fmt(transfer_chiqim) if transfer_chiqim else '—'}</td>
                </tr>
                <tr style="background-color: #fafafa;">
                    <td style="padding: 10px; border: 1px solid #ddd;">Прочие</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #388e3c;">{fmt(other_kirim) if other_kirim else '—'}</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #d32f2f;">{fmt(other_chiqim) if other_chiqim else '—'}</td>
                </tr>
                <tr style="{expense_cursor}" {expense_onclick}>
                    <td style="padding: 10px; border: 1px solid #ddd;">{expense_arrow}Расходы</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #388e3c;">{fmt(expense_kirim) if expense_kirim else '—'}</td>
                    <td style="padding: 10px; border: 1px solid #ddd; text-align: right; color: #d32f2f;">{fmt(expense_chiqim) if expense_chiqim else '—'}</td>
                </tr>
                {expense_sub_rows}
                <tr style="background-color: #e3f2fd; font-weight: bold;">
                    <td style="padding: 12px; border: 1px solid #ddd; font-weight: bold;">Конечный остаток</td>
                    <td style="padding: 12px; border: 1px solid #ddd; text-align: right; font-weight: bold;" colspan="2">{fmt(closing)}</td>
                </tr>
            </tbody>
        </table>
    </div>
    """

    return html


@frappe.whitelist()
def export_dds_excel(filters):
    """DDS'ni Excel qilib qaytaradi: 1-varaq Свод (summary), 2-varaq Транзакции (xom)."""
    import json
    import openpyxl
    from io import BytesIO
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if isinstance(filters, str):
        filters = json.loads(filters)

    data, expense_summaries, employee_group_summaries = get_data(filters)
    s = compute_summary(data, expense_summaries, employee_group_summaries)

    # Свод varaqi — reportdagi tartibda
    summary_rows = [["DDS", "Кирим", "Чиқим"]]
    summary_rows.append(["Начальный остаток", None, s["opening"]])
    for cat in ("customer", "supplier", "dividend", "employee", "transfer", "other"):
        c = s["cats"][cat]
        summary_rows.append([CATEGORY_LABELS[cat], c["kirim"], c["chiqim"]])
        if cat == "employee":
            for group in sorted(s["employee_sub"], key=lambda g: (g == "Без группы", g)):
                t = s["employee_sub"][group]
                summary_rows.append([f"    {group}", t["kirim"], t["chiqim"]])
    exp = s["cats"]["expense"]
    summary_rows.append(["Расходы", exp["kirim"], exp["chiqim"]])
    for desc, t in s["expense_sub"].items():
        name = desc.replace("Расходы: ", "") if desc.startswith("Расходы: ") else desc
        summary_rows.append([f"    {name}", t["kirim"], t["chiqim"]])
    summary_rows.append(["Конечный остаток", None, s["closing"]])

    # Транзакции varaqi — xom qatorlar (Employee Group ustuni Кирим/Чиқим bilan Категория orasida)
    txn_rows = [["Сана", "Касса счёт", "Кирим/Чиқим", "Employee Group", "Категория",
                 "Счёт", "Тип 1",
                 "Сумма", "Валюта", "Изоҳ", "Тип", "Документ"]]
    for r in data:
        txn_rows.append([
            r.get("posting_date"), r.get("account"), r.get("direction"),
            r.get("employee_group"),
            r.get("description"),
            r.get("schet"), r.get("tip1"),
            r.get("summa"), r.get("currency"),
            r.get("remarks"), r.get("voucher_type"), r.get("voucher_no"),
        ])

    def clean(cell):
        if isinstance(cell, str):
            return ILLEGAL_CHARACTERS_RE.sub("", cell)
        return cell

    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Свод"
    for row in summary_rows:
        ws_summary.append([clean(c) for c in row])

    ws_txn = wb.create_sheet("Транзакции")
    for row in txn_rows:
        ws_txn.append([clean(c) for c in row])

    bio = BytesIO()
    wb.save(bio)

    frappe.response["filecontent"] = bio.getvalue()
    frappe.response["filename"] = "DDS.xlsx"
    frappe.response["type"] = "binary"
